"""
Monta a secao "Descricao Tecnica" (memorial descritivo) como subdocumento
docxtpl, inserida no lugar do placeholder {{p descricao_tecnica}} -- apenas
na Proposta Tecnica.

Comeca em pagina nova, com titulo estilo Heading 1 e bookmark _Toc190957976
(reaproveitado no template tecnica para a entrada do Sumario). No fim, uma
linha de "Prazo de execucao" (a Tecnica nao tem a pagina de Prazo e Preco).

A formatacao do corpo replica a do texto corrido do documento original:
Arial Narrow 12pt, cor 1B3462, justificado, com o mesmo recuo.
"""
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

FONT_NAME = "Arial Narrow"
FONT_SIZE = Pt(12)
COR_TEXTO = RGBColor(0x1B, 0x34, 0x62)

# Mesmos recuos do texto corrido do modelo (secao "Escopo"): a pagina tem
# margem 0 no .docx, entao o recuo DIREITO tambem precisa ser explicito senao
# o texto encosta na borda da folha.
RECUO_ESQUERDO = Twips(1701)
RECUO_DIREITO = Twips(1137)
RECUO_PRIMEIRA_LINHA = Twips(459)

BOOKMARK_NAME = "_Toc190957976"
BOOKMARK_ID = 900


def _envolver_com_bookmark(paragraph, nome=BOOKMARK_NAME, bm_id=BOOKMARK_ID):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bm_id))
    start.set(qn("w:name"), nome)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bm_id))

    p = paragraph._p
    pPr = p.find(qn("w:pPr"))
    if pPr is not None:
        pPr.addnext(start)
    else:
        p.insert(0, start)
    p.append(end)


def _paragrafo_corpo(subdoc):
    p = subdoc.add_paragraph()
    p.style = "Normal"
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = RECUO_ESQUERDO
    p.paragraph_format.right_indent = RECUO_DIREITO
    p.paragraph_format.first_line_indent = RECUO_PRIMEIRA_LINHA
    p.paragraph_format.space_after = Pt(6)
    return p


def _formatar(run, bold=False):
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE
    run.font.color.rgb = COR_TEXTO
    run.bold = bold


def montar_secao_descricao_tecnica(subdoc, consideracoes: str = "",
                                   descricao_itens: str = "",
                                   prazo_execucao: str = ""):
    """consideracoes: texto livre de premissas (um paragrafo por linha).
    descricao_itens: texto livre da descricao tecnica dos itens (um paragrafo
    por linha)."""
    titulo = subdoc.add_paragraph()
    try:
        titulo.style = "Heading 1"
    except KeyError:
        pass
    titulo.paragraph_format.page_break_before = True
    titulo.paragraph_format.left_indent = Twips(1701)
    titulo.add_run("Descrição Técnica")
    _envolver_com_bookmark(titulo)

    for linha in (consideracoes or "").split("\n"):
        linha = linha.strip()
        if linha:
            _formatar(_paragrafo_corpo(subdoc).add_run(linha))

    for linha in (descricao_itens or "").split("\n"):
        linha = linha.strip()
        if linha:
            _formatar(_paragrafo_corpo(subdoc).add_run(linha))

    if prazo_execucao:
        _formatar(_paragrafo_corpo(subdoc).add_run(""))  # espacamento
        p = _paragrafo_corpo(subdoc)
        _formatar(p.add_run("Prazo de execução: "), bold=True)
        _formatar(p.add_run(
            f"{prazo_execucao} a partir do aceite da proposta e mobilização "
            "da equipe e materiais."
        ))
